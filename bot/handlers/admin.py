"""Admin handler — mukammal dashboard, batafsil tahliliy statistika,
a'zolar ro'yxati va anketalari, adminlarni boshqarish, qidiruv,
targetli broadcast va Excel (.xlsx) eksport.
"""

import os
import csv
import logging
import asyncio
from datetime import datetime, timedelta

from aiogram import Router, F
from aiogram.filters import Command
from aiogram.types import (
    Message,
    CallbackQuery,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    FSInputFile,
)
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup

from sqlalchemy import select, func, desc, or_, and_, update
from sqlalchemy.ext.asyncio import AsyncSession

from bot.config import settings
from db.session import async_session
from db.models import User, Profile, Referral, InterviewSession
from bot.keyboards.reply import main_menu_keyboard

logger = logging.getLogger(__name__)
router = Router(name="admin")


# ═══════════════════════════════════════════════════════════════════════
# FSM HOLATLARI
# ═══════════════════════════════════════════════════════════════════════

class AdminStates(StatesGroup):
    waiting_for_search_query = State()
    waiting_for_new_admin = State()
    waiting_for_broadcast_target = State()
    waiting_for_broadcast_content = State()
    waiting_for_broadcast_confirm = State()


# ═══════════════════════════════════════════════════════════════════════
# ADMIN HUQUQINI TEKSHIRISH
# ═══════════════════════════════════════════════════════════════════════

async def is_admin_user(telegram_id: int) -> bool:
    """Foydalanuvchi superadmin (.env) yoki bazada admin ekanligini tekshirish."""
    if telegram_id in settings.admin_ids_list:
        return True
    async with async_session() as session:
        user = await session.scalar(select(User).where(User.telegram_id == telegram_id))
        if user and getattr(user, "is_admin", False):
            return True
    return False


def is_superadmin(telegram_id: int) -> bool:
    """Faqat .env dagi asosiy adminlar."""
    return telegram_id in settings.admin_ids_list


# ═══════════════════════════════════════════════════════════════════════
# KLAVIATURALAR
# ═══════════════════════════════════════════════════════════════════════

def admin_main_keyboard() -> InlineKeyboardMarkup:
    """Asosiy admin boshqaruv tugmalari."""
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="📊 Batafsil Statistika", callback_data="admin_stat"),
            InlineKeyboardButton(text="👥 A'zolar ro'yxati", callback_data="admin_users:0:all"),
        ],
        [
            InlineKeyboardButton(text="🔍 A'zoni qidirish", callback_data="admin_search"),
            InlineKeyboardButton(text="👑 Adminlar boshqaruvi", callback_data="admin_manage_admins"),
        ],
        [
            InlineKeyboardButton(text="📥 Excel (.xlsx) yuklash", callback_data="admin_export"),
            InlineKeyboardButton(text="📢 Xabar tarqatish", callback_data="admin_broadcast_menu"),
        ],
        [
            InlineKeyboardButton(text="🔄 Yangilash", callback_data="admin_refresh"),
        ],
    ])


def admin_stat_keyboard() -> InlineKeyboardMarkup:
    """Statistika bo'limidagi ichki tugmalar."""
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="🏢 Sohalar tahlili", callback_data="admin_stat_industries"),
            InlineKeyboardButton(text="🎯 Maqsadlar tahlili", callback_data="admin_stat_goals"),
        ],
        [
            InlineKeyboardButton(text="💼 Tajriba va Yosh", callback_data="admin_stat_experience"),
            InlineKeyboardButton(text="👥 Takliflar (Referral)", callback_data="admin_stat_referrals"),
        ],
        [
            InlineKeyboardButton(text="🔄 Yangilash", callback_data="admin_stat"),
            InlineKeyboardButton(text="⬅️ Bosh menyu", callback_data="admin_home"),
        ],
    ])


def back_to_stat_keyboard() -> InlineKeyboardMarkup:
    """Statistika bo'limiga qaytish tugmasi."""
    return InlineKeyboardMarkup(inline_keyboard=[
        [
            InlineKeyboardButton(text="⬅️ Statistikaga qaytish", callback_data="admin_stat"),
            InlineKeyboardButton(text="🏠 Bosh menyu", callback_data="admin_home"),
        ]
    ])


def back_to_admin_keyboard() -> InlineKeyboardMarkup:
    """Bosh admin menyusiga qaytish."""
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ Bosh menyu", callback_data="admin_home")]
    ])


# ═══════════════════════════════════════════════════════════════════════
# ASOSIY DASHBOARD
# ═══════════════════════════════════════════════════════════════════════

async def _build_dashboard_text() -> str:
    """Dashboard matnini bazadagi eng so'nggi ma'lumotlar bilan tayyorlash."""
    async with async_session() as session:
        # Jami botga kirganlar
        total_users = await session.scalar(select(func.count(User.id))) or 0
        # Telefon raqam berganlar
        users_with_phone = await session.scalar(
            select(func.count(User.id)).where(User.phone.isnot(None))
        ) or 0
        # Anketani to'liq topshirganlar
        completed_profiles = await session.scalar(
            select(func.count(Profile.id)).where(Profile.bio_summary.isnot(None))
        ) or 0
        # Jarayonda (chala) qolganlar
        incomplete_users = max(0, total_users - completed_profiles)

        # Bugungi statistika (oxirgi 24 soat)
        today_start = datetime.utcnow() - timedelta(hours=24)
        today_users = await session.scalar(
            select(func.count(User.id)).where(User.created_at >= today_start)
        ) or 0
        today_completed = await session.scalar(
            select(func.count(Profile.id)).where(
                Profile.created_at >= today_start,
                Profile.bio_summary.isnot(None)
            )
        ) or 0

        # Referral orqali kirganlar
        referral_users = await session.scalar(
            select(func.count(User.id)).where(User.referred_by.isnot(None))
        ) or 0

    # Konversiya darajasi
    conversion_rate = (completed_profiles / total_users * 100) if total_users > 0 else 0.0

    deadline_status = "🟢 Qabul ochiq" if settings.is_registration_open() else "🔴 Qabul yopilgan"

    text = (
        "╔════════════════════════════╗\n"
        "   ☕ TEAHOUSE ADMIN DASHBOARD   \n"
        "╚════════════════════════════╝\n\n"
        "📊 **ASOSIY METRIKALAR:**\n"
        f"• 👥 Jami botga kirganlar: {total_users:,} nafar\n"
        f"• 📱 Telefon ulashganlar: {users_with_phone:,} nafar\n"
        f"• ✅ Anketani to'liq topshirganlar: {completed_profiles:,} nafar\n"
        f"• ⏳ Anketasi chala qolganlar: {incomplete_users:,} nafar\n"
        f"• 📈 Umumiy konversiya: {conversion_rate:.1f}%\n"
        f"• 🔗 Do'sti taklifi bilan kelgan: {referral_users:,} nafar\n\n"
        "📅 **OXIRGI 24 SOAT (BUGUN):**\n"
        f"• 🆕 Yangi kirganlar: +{today_users} kishi\n"
        f"• 📝 Yangi topshirilgan anketalar: +{today_completed} ta\n\n"
        "⏰ **QABUL VA DEADLAYN:**\n"
        f"• Holat: {deadline_status}\n"
        f"• Muddat: {settings.registration_deadline}\n\n"
        "Kerakli bo'limni tanlash uchun quyidagi tugmalardan foydalaning:"
    )
    return text


@router.message(Command("admin"))
@router.message(F.text.in_(["⚙️ Admin boshqaruvi", "Admin boshqaruvi", "Admin"]))
async def admin_dashboard_handler(message: Message, state: FSMContext):
    """Admin boshqaruv markazini ochish."""
    await state.clear()
    if not await is_admin_user(message.from_user.id):
        await message.answer(
            "⛔ <b>Kechirasiz, sizda adminlik huquqi yo'q.</b>\n\n"
            f"Sizning Telegram ID: <code>{message.from_user.id}</code>\n"
            "Ushbu ID ni bosh adminga yuborib, ruxsat olishingiz mumkin.",
            parse_mode="HTML"
        )
        return

    text = await _build_dashboard_text()
    await message.answer(
        text,
        reply_markup=admin_main_keyboard(),
        parse_mode="Markdown"
    )


@router.callback_query(F.data.in_(["admin_home", "admin_refresh"]))
async def admin_home_callback(callback: CallbackQuery, state: FSMContext):
    """Bosh sahifaga qaytish yoki yangilash callback."""
    await state.clear()
    if not await is_admin_user(callback.from_user.id):
        await callback.answer("Ruxsat berilmagan!", show_alert=True)
        return

    await callback.answer("Yangilandi ✅")
    text = await _build_dashboard_text()
    try:
        await callback.message.edit_text(
            text,
            reply_markup=admin_main_keyboard(),
            parse_mode="Markdown"
        )
    except Exception:
        await callback.message.answer(
            text,
            reply_markup=admin_main_keyboard(),
            parse_mode="Markdown"
        )


# ═══════════════════════════════════════════════════════════════════════
# TAHLILIY STATISTIKA (ANALYTICS)
# ═══════════════════════════════════════════════════════════════════════

@router.message(Command("stat"))
@router.callback_query(F.data == "admin_stat")
async def admin_stat_overview(event: Message | CallbackQuery, state: FSMContext):
    """Umumiy voronka va vaqtinchalik o'sish statistikasi."""
    await state.clear()
    user_id = event.from_user.id
    if not await is_admin_user(user_id):
        if isinstance(event, CallbackQuery):
            await event.answer("Ruxsat berilmagan!", show_alert=True)
        return

    if isinstance(event, CallbackQuery):
        await event.answer()
        reply_target = event.message
    else:
        reply_target = event

    async with async_session() as session:
        total = await session.scalar(select(func.count(User.id))) or 0
        phones = await session.scalar(
            select(func.count(User.id)).where(User.phone.isnot(None))
        ) or 0
        completed = await session.scalar(
            select(func.count(Profile.id)).where(Profile.bio_summary.isnot(None))
        ) or 0

        # Dinamika: 24 soat, 7 kun, 30 kun
        now = datetime.utcnow()
        day1 = now - timedelta(days=1)
        day7 = now - timedelta(days=7)
        day30 = now - timedelta(days=30)

        u_24h = await session.scalar(select(func.count(User.id)).where(User.created_at >= day1)) or 0
        c_24h = await session.scalar(
            select(func.count(Profile.id)).where(Profile.created_at >= day1, Profile.bio_summary.isnot(None))
        ) or 0

        u_7d = await session.scalar(select(func.count(User.id)).where(User.created_at >= day7)) or 0
        c_7d = await session.scalar(
            select(func.count(Profile.id)).where(Profile.created_at >= day7, Profile.bio_summary.isnot(None))
        ) or 0

        u_30d = await session.scalar(select(func.count(User.id)).where(User.created_at >= day30)) or 0
        c_30d = await session.scalar(
            select(func.count(Profile.id)).where(Profile.created_at >= day30, Profile.bio_summary.isnot(None))
        ) or 0

    p_phone = (phones / total * 100) if total > 0 else 0
    p_comp = (completed / total * 100) if total > 0 else 0
    p_drop = (100 - p_comp) if total > 0 else 0

    text = (
        "📊 **TEAHOUSE CHUQUR TAHLILIY STATISTIKASI**\n\n"
        "🎯 **KONVERSIYA VORONKASI (FUNNEL):**\n"
        f"1️⃣ Botga kirganlar: **{total:,}** (100%)\n"
        f"   │\n"
        f"   ▼\n"
        f"2️⃣ Telefon qoldirganlar: **{phones:,}** ({p_phone:.1f}%)\n"
        f"   │\n"
        f"   ▼\n"
        f"3️⃣ Anketani to'liq topshirganlar: **{completed:,}** ({p_comp:.1f}%)\n"
        f"⚠️ Jarayonda to'xtaganlar (Chala): **{total - completed:,}** ({p_drop:.1f}%)\n\n"
        "📈 **FOYDALANUVCHILAR O'SISH DINAMIKASI:**\n"
        f"• Bugun (24 soat): **+{u_24h}** kirdi, **+{c_24h}** anketa\n"
        f"• Oxirgi 7 kun: **+{u_7d}** kirdi, **+{c_7d}** anketa\n"
        f"• Oxirgi 30 kun: **+{u_30d}** kirdi, **+{c_30d}** anketa\n\n"
        "Quyidagi tugmalar orqali sohalar, maqsadlar, yosh va takliflar bo'yicha tahlillarni ko'rishingiz mumkin:"
    )

    if isinstance(event, CallbackQuery):
        await reply_target.edit_text(text, reply_markup=admin_stat_keyboard(), parse_mode="Markdown")
    else:
        await reply_target.answer(text, reply_markup=admin_stat_keyboard(), parse_mode="Markdown")


@router.callback_query(F.data == "admin_stat_industries")
async def admin_stat_industries_callback(callback: CallbackQuery):
    """Sohalar bo'yicha taqsimot."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    async with async_session() as session:
        query = (
            select(Profile.industry, func.count(Profile.id))
            .where(Profile.bio_summary.isnot(None), Profile.industry.isnot(None))
            .group_by(Profile.industry)
            .order_by(desc(func.count(Profile.id)))
        )
        rows = (await session.execute(query)).all()
        total_p = sum(r[1] for r in rows) or 1

    if not rows:
        await callback.message.edit_text(
            "🏢 **Sohalar bo'yicha taqsimot:**\n\nHozircha to'ldirilgan anketalar mavjud emas.",
            reply_markup=back_to_stat_keyboard(),
            parse_mode="Markdown"
        )
        return

    lines = ["🏢 **FAOLIYAT SOHALARI BO'YICHA TAQSIMOT:**\n"]
    for idx, (ind, cnt) in enumerate(rows, 1):
        pct = (cnt / total_p) * 100
        lines.append(f"{idx}. **{ind}**: {cnt} kishi ({pct:.1f}%)")

    lines.append(f"\nJami ko'rib chiqilgan anketalar: **{total_p}** ta")
    await callback.message.edit_text("\n".join(lines), reply_markup=back_to_stat_keyboard(), parse_mode="Markdown")


@router.callback_query(F.data == "admin_stat_goals")
async def admin_stat_goals_callback(callback: CallbackQuery):
    """Uchrashuvdan ko'zlangan maqsadlar bo'yicha tahlil."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    async with async_session() as session:
        query = (
            select(Profile.target_partner, func.count(Profile.id))
            .where(Profile.bio_summary.isnot(None), Profile.target_partner.isnot(None))
            .group_by(Profile.target_partner)
            .order_by(desc(func.count(Profile.id)))
        )
        rows = (await session.execute(query)).all()
        total_p = sum(r[1] for r in rows) or 1

    if not rows:
        await callback.message.edit_text(
            "🎯 **Maqsadlar bo'yicha taqsimot:**\n\nHozircha ma'lumotlar yo'q.",
            reply_markup=back_to_stat_keyboard(),
            parse_mode="Markdown"
        )
        return

    lines = ["🎯 **QIDIRILAYOTGAN SHERIK / MAQSADLAR:**\n"]
    for idx, (goal, cnt) in enumerate(rows, 1):
        pct = (cnt / total_p) * 100
        lines.append(f"{idx}. **{goal}**: {cnt} kishi ({pct:.1f}%)")

    lines.append(f"\nJami ko'rib chiqilgan anketalar: **{total_p}** ta")
    await callback.message.edit_text("\n".join(lines), reply_markup=back_to_stat_keyboard(), parse_mode="Markdown")


@router.callback_query(F.data == "admin_stat_experience")
async def admin_stat_experience_callback(callback: CallbackQuery):
    """Tajriba darajasi va yosh toifalari bo'yicha tahlil."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    async with async_session() as session:
        # Tajriba darajasi
        sen_query = (
            select(Profile.seniority, func.count(Profile.id))
            .where(Profile.bio_summary.isnot(None), Profile.seniority.isnot(None))
            .group_by(Profile.seniority)
            .order_by(desc(func.count(Profile.id)))
        )
        sen_rows = (await session.execute(sen_query)).all()

        # Yosh toifalari
        age_query = (
            select(Profile.age, func.count(Profile.id))
            .where(Profile.bio_summary.isnot(None), Profile.age.isnot(None))
            .group_by(Profile.age)
            .order_by(desc(func.count(Profile.id)))
        )
        age_rows = (await session.execute(age_query)).all()

    lines = ["💼 **TAJRIBA DARAJALARI BO'YICHA:**\n"]
    if sen_rows:
        for idx, (sen, cnt) in enumerate(sen_rows, 1):
            lines.append(f"• **{sen}**: {cnt} kishi")
    else:
        lines.append("Hozircha ma'lumot yo'q")

    lines.append("\n🎂 **YOSH GURUHLARI BO'YICHA:**\n")
    if age_rows:
        for idx, (age, cnt) in enumerate(age_rows, 1):
            lines.append(f"• Yosh **{age}**: {cnt} kishi")
    else:
        lines.append("Hozircha ma'lumot yo'q")

    await callback.message.edit_text("\n".join(lines), reply_markup=back_to_stat_keyboard(), parse_mode="Markdown")


@router.callback_query(F.data == "admin_stat_referrals")
async def admin_stat_referrals_callback(callback: CallbackQuery):
    """Referral (do'stlarni taklif qilish) tizimi tahlili."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    async with async_session() as session:
        total_refs = await session.scalar(
            select(func.count(Referral.id))
        ) or 0

        # Eng ko'p taklif qilgan 5 kishi
        top_query = (
            select(User.id, User.first_name, User.username, func.count(Referral.id).label("ref_cnt"))
            .join(Referral, User.id == Referral.referrer_id)
            .group_by(User.id, User.first_name, User.username)
            .order_by(desc("ref_cnt"))
            .limit(5)
        )
        top_rows = (await session.execute(top_query)).all()

    lines = [
        "👥 **REFERRAL VA DO'STLARNI TAKLIF QILISH TAHLILI:**\n",
        f"🔗 Jami taklif havolalari orqali kirganlar: **{total_refs}** nafar\n",
        "🏆 **ENG FAOL TAKLIF QILUVCHILAR (TOP 5):**\n"
    ]

    if top_rows:
        for idx, (uid, fname, uname, cnt) in enumerate(top_rows, 1):
            user_handle = f"@{uname}" if uname else f"ID: {uid}"
            lines.append(f"{idx}. **{fname}** ({user_handle}) — **{cnt} ta** taklif")
    else:
        lines.append("Hozircha taklif orqali kirganlar mavjud emas.")

    await callback.message.edit_text("\n".join(lines), reply_markup=back_to_stat_keyboard(), parse_mode="Markdown")


# ═══════════════════════════════════════════════════════════════════════
# A'ZOLAR RO'YXATI VA HAR BIR A'ZO ANKETASINI KO'RISH
# ═══════════════════════════════════════════════════════════════════════

PAGE_SIZE = 5

async def _get_users_page(page: int, filter_type: str):
    """Sahifalangan foydalanuvchilar ro'yxati va umumiy sahifalar soni."""
    async with async_session() as session:
        base_query = select(User, Profile).outerjoin(Profile, User.id == Profile.user_id)
        count_query = select(func.count(User.id)).outerjoin(Profile, User.id == Profile.user_id)

        if filter_type == "completed":
            base_query = base_query.where(Profile.bio_summary.isnot(None))
            count_query = count_query.where(Profile.bio_summary.isnot(None))
        elif filter_type == "incomplete":
            base_query = base_query.where(or_(Profile.id.is_(None), Profile.bio_summary.is_(None)))
            count_query = count_query.where(or_(Profile.id.is_(None), Profile.bio_summary.is_(None)))

        total_count = await session.scalar(count_query) or 0
        total_pages = max(1, (total_count + PAGE_SIZE - 1) // PAGE_SIZE)

        page = max(0, min(page, total_pages - 1))

        query = base_query.order_by(desc(User.created_at)).offset(page * PAGE_SIZE).limit(PAGE_SIZE)
        rows = (await session.execute(query)).all()

        return rows, total_count, total_pages, page


def _build_users_keyboard(rows, page: int, total_pages: int, filter_type: str) -> InlineKeyboardMarkup:
    """A'zolar ro'yxati tugmalari (filtrlar, a'zolar, pagination)."""
    buttons = []

    # Filtrlar
    f_all = "🔘 Barchasi" if filter_type == "all" else "Barchasi"
    f_comp = "✅ To'liq" if filter_type == "completed" else "To'liq"
    f_inc = "⏳ Chala" if filter_type == "incomplete" else "Chala"

    buttons.append([
        InlineKeyboardButton(text=f_all, callback_data=f"admin_users:0:all"),
        InlineKeyboardButton(text=f_comp, callback_data=f"admin_users:0:completed"),
        InlineKeyboardButton(text=f_inc, callback_data=f"admin_users:0:incomplete"),
    ])

    # Foydalanuvchilar qatori
    for u, p in rows:
        status_icon = "✅" if (p and p.bio_summary) else ("📱" if u.phone else "⏳")
        role_preview = (p.role[:14] + "…") if (p and p.role and len(p.role) > 14) else (p.role if p and p.role else "Anketasiz")
        btn_text = f"{status_icon} {u.first_name[:15]} ({role_preview})"
        buttons.append([
            InlineKeyboardButton(text=btn_text, callback_data=f"admin_uview:{u.id}:{page}:{filter_type}")
        ])

    # Sahifalash (Pagination)
    nav_row = []
    if page > 0:
        nav_row.append(InlineKeyboardButton(text="◀️ Oldingi", callback_data=f"admin_users:{page - 1}:{filter_type}"))
    nav_row.append(InlineKeyboardButton(text=f"📄 {page + 1}/{total_pages}", callback_data="admin_noop"))
    if page < total_pages - 1:
        nav_row.append(InlineKeyboardButton(text="Keyingi ▶️", callback_data=f"admin_users:{page + 1}:{filter_type}"))

    if nav_row:
        buttons.append(nav_row)

    # Bosh menyu
    buttons.append([InlineKeyboardButton(text="⬅️ Bosh menyu", callback_data="admin_home")])

    return InlineKeyboardMarkup(inline_keyboard=buttons)


@router.message(Command("users"))
@router.callback_query(F.data.startswith("admin_users:"))
async def admin_users_list_handler(event: Message | CallbackQuery):
    """A'zolar ro'yxati (filtrlash va sahifalash bilan)."""
    user_id = event.from_user.id
    if not await is_admin_user(user_id):
        return

    if isinstance(event, CallbackQuery):
        await event.answer()
        parts = event.data.split(":")
        page = int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else 0
        filter_type = parts[2] if len(parts) > 2 else "all"
        reply_target = event.message
    else:
        page = 0
        filter_type = "all"
        reply_target = event

    rows, total_count, total_pages, page = await _get_users_page(page, filter_type)

    filter_names = {"all": "Barcha foydalanuvchilar", "completed": "Faqat to'liq anketalar", "incomplete": "Faqat chala anketalar"}
    header_text = (
        f"👥 **A'ZOLAR RO'YXATI** ({filter_names.get(filter_type, 'Barchasi')})\n\n"
        f"Jami a'zolar: **{total_count} nafar**\n"
        "Batafsil anketasini ko'rish uchun quyidagi a'zolardan birini tanlang:\n"
    )

    kb = _build_users_keyboard(rows, page, total_pages, filter_type)
    if isinstance(event, CallbackQuery):
        await reply_target.edit_text(header_text, reply_markup=kb, parse_mode="Markdown")
    else:
        await reply_target.answer(header_text, reply_markup=kb, parse_mode="Markdown")


@router.callback_query(F.data == "admin_noop")
async def admin_noop_callback(callback: CallbackQuery):
    """Sahifa raqami bosilganda hech narsa qilmaslik."""
    await callback.answer()


@router.callback_query(F.data.startswith("admin_uview:"))
async def admin_user_view_callback(callback: CallbackQuery):
    """A'zoning barcha 13 ta savoldan iborat to'liq anketasini ochish."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    parts = callback.data.split(":")
    target_user_id = int(parts[1])
    page = int(parts[2]) if len(parts) > 2 else 0
    filter_type = parts[3] if len(parts) > 3 else "all"

    async with async_session() as session:
        query = (
            select(User, Profile)
            .outerjoin(Profile, User.id == Profile.user_id)
            .where(User.id == target_user_id)
        )
        res = (await session.execute(query)).first()

    if not res:
        await callback.message.edit_text(
            "Foydalanuvchi topilmadi.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="⬅️ Ortga", callback_data=f"admin_users:{page}:{filter_type}")]
            ])
        )
        return

    u, p = res
    status_str = "✅ To'liq topshirgan" if (p and p.bio_summary) else ("📱 Telefon qoldirgan" if u.phone else "⏳ Botga kirgan")
    admin_str = "👑 Admin" if getattr(u, "is_admin", False) or is_superadmin(u.telegram_id) else "Foydalanuvchi"
    created_str = u.created_at.strftime("%Y-%m-%d %H:%M") if hasattr(u, "created_at") and u.created_at else "Noma'lum"

    card = (
        "╔════════════════════════════╗\n"
        f"   👤 A'ZO ANKETASI: {u.first_name}   \n"
        "╚════════════════════════════╝\n\n"
        "📋 **SHAXSIY MA'LUMOTLAR:**\n"
        f"• Telegram ID: `{u.telegram_id}`\n"
        f"• Username: @{u.username or 'yoqd'}\n"
        f"• Telefon: {u.phone or 'kiritilmagan'}\n"
        f"• Holat: **{status_str}**\n"
        f"• Tizimdagi maqomi: **{admin_str}**\n"
        f"• Kirgan vaqti: {created_str}\n\n"
    )

    if p and p.bio_summary:
        card += (
            "💼 **1-QISM: KASBI VA NATIJALARI:**\n"
            f"• Kasbi va lavozimi: {p.role or '—'}\n"
            f"• Kompaniya / Loyiha: {p.company or '—'}\n"
            f"• Faoliyat sohasi: {p.industry or '—'}\n"
            f"• Tajribasi: {p.seniority or '—'} ({p.experience_years or 0} yil)\n"
            f"• Erishgan eng katta yutug'i:\n  _{p.achievements or '—'}_\n"
            f"• Yoshi: {p.age or '—'}\n\n"
            "🤝 **2-QISM: QIDIRILAYOTGAN SHERIKLAR:**\n"
            f"• Uchrashuvdan maqsadi: {p.target_partner or p.current_goal or '—'}\n"
            f"• Qiziqqan soha vakillari: {p.target_industry or 'Barchasi'}\n"
            f"• Sherikning tajriba darajasi: {p.target_seniority or '—'}\n"
            f"• O'zi nima taklif qila oladi:\n  _{p.can_offer or '—'}_\n"
            f"• Muhokama mavzusi:\n  _{p.interests or '—'}_\n\n"
            f"🤖 **AI BIO XULOSASI:**\n\"{p.bio_summary}\"\n"
        )
    else:
        card += "⚠️ **Eslatma:** Ushbu a'zo hali anketani to'liq yakunlamagan.\n"

    # Tugmalar
    buttons = []
    is_target_admin = getattr(u, "is_admin", False) or is_superadmin(u.telegram_id)

    # Admin berish / olish tugmasi
    if not is_target_admin:
        buttons.append([
            InlineKeyboardButton(text="👑 Admin huquqini berish", callback_data=f"admin_grant:{u.id}:{page}:{filter_type}")
        ])
    elif not is_superadmin(u.telegram_id):
        buttons.append([
            InlineKeyboardButton(text="❌ Adminlikdan olish", callback_data=f"admin_revoke:{u.id}:{page}:{filter_type}")
        ])

    # Profilga havola
    if u.username:
        buttons.append([
            InlineKeyboardButton(text="💬 Telegram orqali yozish", url=f"https://t.me/{u.username}")
        ])

    buttons.append([
        InlineKeyboardButton(text="⬅️ Ro'yxatga qaytish", callback_data=f"admin_users:{page}:{filter_type}")
    ])

    await callback.message.edit_text(card, reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons), parse_mode="Markdown")


# ═══════════════════════════════════════════════════════════════════════
# ADMIN HUQUQLARINI BOSHQARISH (ADMIN TAYINLASH VA OLISH)
# ═══════════════════════════════════════════════════════════════════════

@router.callback_query(F.data.startswith("admin_grant:"))
async def admin_grant_callback(callback: CallbackQuery):
    """Foydalanuvchiga adminlik huquqini berish."""
    if not await is_admin_user(callback.from_user.id):
        return

    parts = callback.data.split(":")
    target_id = int(parts[1])
    page = int(parts[2]) if len(parts) > 2 else 0
    filter_type = parts[3] if len(parts) > 3 else "all"

    async with async_session() as session:
        user = await session.scalar(select(User).where(User.id == target_id))
        if user:
            user.is_admin = True
            await session.commit()
            await callback.answer(f"✅ {user.first_name} ga admin huquqi berildi!", show_alert=True)
            try:
                await callback.bot.send_message(
                    chat_id=user.telegram_id,
                    text="🎉 **Tabriklaymiz!** Sizga Teahouse botida adminlik huquqi berildi.\n\n/admin buyrug'i orqali boshqaruv markazini ochishingiz mumkin.",
                    reply_markup=main_menu_keyboard(is_admin=True),
                    parse_mode="Markdown"
                )
            except Exception:
                pass
        else:
            await callback.answer("Foydalanuvchi topilmadi!", show_alert=True)

    callback.data = f"admin_uview:{target_id}:{page}:{filter_type}"
    await admin_user_view_callback(callback)


@router.callback_query(F.data.startswith("admin_revoke:"))
async def admin_revoke_callback(callback: CallbackQuery):
    """Foydalanuvchidan adminlik huquqini bekor qilish."""
    if not await is_admin_user(callback.from_user.id):
        return

    parts = callback.data.split(":")
    target_id = int(parts[1])
    page = int(parts[2]) if len(parts) > 2 else 0
    filter_type = parts[3] if len(parts) > 3 else "all"

    async with async_session() as session:
        user = await session.scalar(select(User).where(User.id == target_id))
        if user:
            if is_superadmin(user.telegram_id):
                await callback.answer("⚠️ Bosh admin (Superadmin) huquqini bekor qilib bo'lmaydi!", show_alert=True)
                return
            user.is_admin = False
            await session.commit()
            await callback.answer(f"❌ {user.first_name} adminlikdan chiqarildi!", show_alert=True)
            try:
                await callback.bot.send_message(
                    chat_id=user.telegram_id,
                    text="ℹ️ Sizning Teahouse botidagi adminlik huquqingiz to'xtatildi.",
                    reply_markup=main_menu_keyboard(is_admin=False),
                    parse_mode="Markdown"
                )
            except Exception:
                pass

    callback.data = f"admin_uview:{target_id}:{page}:{filter_type}"
    await admin_user_view_callback(callback)


@router.callback_query(F.data == "admin_manage_admins")
async def admin_manage_admins_overview(callback: CallbackQuery):
    """Barcha adminlar ro'yxati va yangi admin tayinlash menyusi."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    async with async_session() as session:
        db_admins = (await session.execute(
            select(User).where(User.is_admin.is_(True))
        )).scalars().all()
        super_ids = settings.admin_ids_list

    lines = [
        "👑 **BOT ADMINLARINI BOSHQARISH**\n\n"
        "⭐️ **Superadminlar (.env orqali doimiy himoyalangan):**"
    ]
    for sid in super_ids:
        lines.append(f"• Telegram ID: `{sid}`")

    lines.append("\n👥 **Tayinlangan adminlar (Baza orqali):**")
    buttons = []
    if db_admins:
        for a in db_admins:
            uname = f"@{a.username}" if a.username else f"ID: {a.telegram_id}"
            lines.append(f"• **{a.first_name}** ({uname})")
            buttons.append([
                InlineKeyboardButton(text=f"❌ {a.first_name} (Adminlikdan olish)", callback_data=f"admin_revoke:{a.id}:0:all")
            ])
    else:
        lines.append("Hozircha qo'shimcha tayinlangan adminlar yo'q.")

    lines.append("\nYangi admin tayinlash uchun pastdagi tugmani bosing:")

    buttons.insert(0, [InlineKeyboardButton(text="➕ Yangi admin qo'shish", callback_data="admin_add_admin_prompt")])
    buttons.append([InlineKeyboardButton(text="⬅️ Bosh menyu", callback_data="admin_home")])

    await callback.message.edit_text(
        "\n".join(lines),
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="Markdown"
    )


@router.callback_query(F.data == "admin_add_admin_prompt")
async def admin_add_admin_prompt(callback: CallbackQuery, state: FSMContext):
    """Yangi admin qo'shish uchun ma'lumot so'rash."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    await state.set_state(AdminStates.waiting_for_new_admin)
    await callback.message.edit_text(
        "👑 **YANGI ADMIN TAYINLASH**\n\n"
        "Admin qilmoqchi bo'lgan foydalanuvchining **Telegram ID raqami** yoki **@username** sini yuboring:\n\n"
        "_(Foydalanuvchi botga kamida bir marta /start bosgan bo'lishi kerak)_\n\n"
        "Bekor qilish uchun /cancel deb yozing.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ Bekor qilish", callback_data="admin_manage_admins")]
        ]),
        parse_mode="Markdown"
    )


@router.message(AdminStates.waiting_for_new_admin, F.text)
async def process_new_admin_input(message: Message, state: FSMContext):
    """Kiritilgan ma'lumot bo'yicha foydalanuvchini topib, admin qilish."""
    if not await is_admin_user(message.from_user.id):
        return

    raw_input = message.text.strip()
    if raw_input.lower() == "/cancel":
        await state.clear()
        await message.answer("Bekor qilindi.", reply_markup=admin_main_keyboard())
        return

    clean_input = raw_input.replace("@", "")

    async with async_session() as session:
        if clean_input.isdigit():
            user = await session.scalar(select(User).where(User.telegram_id == int(clean_input)))
        else:
            user = await session.scalar(select(User).where(func.lower(User.username) == clean_input.lower()))

        if not user:
            await message.answer(
                f"❌ Foydalanuvchi `{raw_input}` bazadan topilmadi.\n\n"
                "Ushbu shaxs botga /start bosganligiga ishonch hosil qiling va qayta urinib ko'ring (yoki /cancel yozing).",
                parse_mode="Markdown"
            )
            return

        user.is_admin = True
        await session.commit()

        try:
            await message.bot.send_message(
                chat_id=user.telegram_id,
                text="🎉 **Tabriklaymiz!** Sizga Teahouse botida adminlik huquqi berildi.\n\n/admin buyrug'i orqali boshqaruv markaziga kirishingiz mumkin.",
                reply_markup=main_menu_keyboard(is_admin=True),
                parse_mode="Markdown"
            )
        except Exception:
            pass

    await state.clear()
    await message.answer(
        f"✅ **Muvaffaqiyatli!**\n\n"
        f"Foydalanuvchi: **{user.first_name}** (@{user.username or 'yoqd'})\n"
        f"Telegram ID: `{user.telegram_id}` endi bot admini hisoblanadi.",
        reply_markup=admin_main_keyboard(),
        parse_mode="Markdown"
    )


# ═══════════════════════════════════════════════════════════════════════
# A'ZONI QIDIRISH (SEARCH)
# ═══════════════════════════════════════════════════════════════════════

@router.callback_query(F.data == "admin_search")
async def admin_search_prompt(callback: CallbackQuery, state: FSMContext):
    """Qidiruv so'zini kiritishni so'rash."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    await state.set_state(AdminStates.waiting_for_search_query)
    await callback.message.edit_text(
        "🔍 **A'ZOLARNI QIDIRISH**\n\n"
        "Qidirilayotgan shaxsning **Ismi**, **@username**, **Telefon raqami** yoki **Telegram ID** sini yuboring:\n\n"
        "Bekor qilish uchun pastdagi tugmani bosing.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ Bekor qilish", callback_data="admin_home")]
        ]),
        parse_mode="Markdown"
    )


@router.message(AdminStates.waiting_for_search_query, F.text)
async def process_search_query(message: Message, state: FSMContext):
    """Qidiruv natijalarini chiqarish."""
    if not await is_admin_user(message.from_user.id):
        return

    query_str = message.text.strip()
    if query_str.lower() == "/cancel":
        await state.clear()
        await message.answer("Qidiruv bekor qilindi.", reply_markup=admin_main_keyboard())
        return

    await state.clear()
    async with async_session() as session:
        conditions = [
            User.first_name.ilike(f"%{query_str}%"),
            User.username.ilike(f"%{query_str.replace('@', '')}%"),
            User.phone.ilike(f"%{query_str}%"),
        ]
        if query_str.isdigit():
            conditions.append(User.telegram_id == int(query_str))

        search_query = (
            select(User, Profile)
            .outerjoin(Profile, User.id == Profile.user_id)
            .where(or_(*conditions))
            .limit(10)
        )
        results = (await session.execute(search_query)).all()

    if not results:
        await message.answer(
            f"🔍 '{query_str}' so'rovi bo'yicha hech qanday a'zo topilmadi.",
            reply_markup=InlineKeyboardMarkup(inline_keyboard=[
                [InlineKeyboardButton(text="🔍 Qayta qidirish", callback_data="admin_search")],
                [InlineKeyboardButton(text="⬅️ Bosh menyu", callback_data="admin_home")],
            ])
        )
        return

    buttons = []
    lines = [f"🔍 **QIDIRUV NATIJALARI ('{query_str}'):**\n"]
    for idx, (u, p) in enumerate(results, 1):
        status = "✅" if (p and p.bio_summary) else ("📱" if u.phone else "⏳")
        uname = f"@{u.username}" if u.username else "username yo'q"
        lines.append(f"{idx}. {status} **{u.first_name}** ({uname}) | Tel: {u.phone or 'yo\'q'}")
        buttons.append([
            InlineKeyboardButton(text=f"👁 {u.first_name} anketasini ko'rish", callback_data=f"admin_uview:{u.id}:0:all")
        ])

    buttons.append([
        InlineKeyboardButton(text="🔍 Yangi qidiruv", callback_data="admin_search"),
        InlineKeyboardButton(text="⬅️ Bosh menyu", callback_data="admin_home"),
    ])

    await message.answer(
        "\n".join(lines),
        reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons),
        parse_mode="Markdown"
    )


# ═══════════════════════════════════════════════════════════════════════
# TARGETLI XABARNOMA (BROADCAST)
# ═══════════════════════════════════════════════════════════════════════

@router.message(Command("broadcast"))
@router.callback_query(F.data == "admin_broadcast_menu")
async def admin_broadcast_menu_handler(event: Message | CallbackQuery, state: FSMContext):
    """Target auditoriyani tanlash menyusi."""
    await state.clear()
    user_id = event.from_user.id
    if not await is_admin_user(user_id):
        return

    if isinstance(event, CallbackQuery):
        await event.answer()
        reply_target = event.message
    else:
        reply_target = event

    async with async_session() as session:
        total = await session.scalar(select(func.count(User.id))) or 0
        comp = await session.scalar(select(func.count(Profile.id)).where(Profile.bio_summary.isnot(None))) or 0
        inc = max(0, total - comp)

    kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"📢 Barchaga ({total} kishi)", callback_data="admin_bctarget:all")],
        [InlineKeyboardButton(text=f"✅ Faqat to'liq anketalarga ({comp} kishi)", callback_data="admin_bctarget:completed")],
        [InlineKeyboardButton(text=f"⏳ Faqat chala qoldirganlarga ({inc} kishi)", callback_data="admin_bctarget:incomplete")],
        [InlineKeyboardButton(text="⬅️ Bosh menyu", callback_data="admin_home")],
    ])

    text = (
        "📢 **XABARNOMA YUBORISH (TARGETED BROADCAST)**\n\n"
        "Xabarni qaysi auditoriyaga yubormoqchisiz?\n\n"
        f"• **Barchaga**: Botga kirgan barcha a'zolarga ({total} kishi)\n"
        f"• **Faqat to'liq anketalarga**: Saralashdan o'tayotganlarga ({comp} kishi)\n"
        f"• **Faqat chala qoldirganlarga**: Anketani to'ldirishni eslatish uchun ({inc} kishi)"
    )

    if isinstance(event, CallbackQuery):
        await reply_target.edit_text(text, reply_markup=kb, parse_mode="Markdown")
    else:
        await reply_target.answer(text, reply_markup=kb, parse_mode="Markdown")


@router.callback_query(F.data.startswith("admin_bctarget:"))
async def admin_bctarget_callback(callback: CallbackQuery, state: FSMContext):
    """Auditoriya tanlangach, xabar matnini so'rash."""
    if not await is_admin_user(callback.from_user.id):
        return

    await callback.answer()
    target_group = callback.data.split(":")[1]
    await state.update_data(target_group=target_group)
    await state.set_state(AdminStates.waiting_for_broadcast_content)

    group_labels = {
        "all": "Barcha botga kirganlarga",
        "completed": "Faqat anketani to'liq topshirganlarga",
        "incomplete": "Faqat anketani to'ldirmaganlarga"
    }

    await callback.message.edit_text(
        f"📢 Auditoriya tanlandi: **{group_labels.get(target_group)}**\n\n"
        "Endi yuboriladigan xabarni yozing (matn, rasm yoki havola bo'lishi mumkin):\n\n"
        "Bekor qilish uchun pastdagi tugmani bosing.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="❌ Bekor qilish", callback_data="admin_broadcast_menu")]
        ]),
        parse_mode="Markdown"
    )


@router.message(AdminStates.waiting_for_broadcast_content)
async def process_broadcast_content(message: Message, state: FSMContext):
    """Xabarni qabul qilish va tasdiqlash oynasini ko'rsatish."""
    if not await is_admin_user(message.from_user.id):
        return

    data = await state.get_data()
    target_group = data.get("target_group", "all")

    text = message.text or message.caption or ""
    photo_id = message.photo[-1].file_id if message.photo else None

    if not text and not photo_id:
        await message.answer("Iltimos, matn yoki rasm yuboring.")
        return

    await state.update_data(text=text, photo_id=photo_id)
    await state.set_state(AdminStates.waiting_for_broadcast_confirm)

    async with async_session() as session:
        if target_group == "completed":
            count = await session.scalar(
                select(func.count(Profile.id)).where(Profile.bio_summary.isnot(None))
            ) or 0
        elif target_group == "incomplete":
            total = await session.scalar(select(func.count(User.id))) or 0
            comp = await session.scalar(select(func.count(Profile.id)).where(Profile.bio_summary.isnot(None))) or 0
            count = max(0, total - comp)
        else:
            count = await session.scalar(select(func.count(User.id))) or 0

    confirm_kb = InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="🚀 Ha, darhol yuborilsin!", callback_data="admin_bc_send_now")],
        [InlineKeyboardButton(text="❌ Bekor qilish", callback_data="admin_broadcast_menu")],
    ])

    await message.answer(
        "⚠️ **XABARNI YUBORISHNI TASDIQLAYSIZMI?**\n\n"
        f"Auditoriya: **{target_group}**\n"
        f"Qabul qiluvchilar soni: **{count} nafar**\n\n"
        "Yuqoridagi xabaringiz to'g'ri bo'lsa, 'Ha, darhol yuborilsin' tugmasini bosing:",
        reply_markup=confirm_kb,
        parse_mode="Markdown"
    )


@router.callback_query(F.data == "admin_bc_send_now")
async def execute_broadcast(callback: CallbackQuery, state: FSMContext):
    """Xabarni haqiqatda tarqatish."""
    if not await is_admin_user(callback.from_user.id):
        return

    data = await state.get_data()
    await state.clear()
    await callback.answer()

    target_group = data.get("target_group", "all")
    text = data.get("text", "")
    photo_id = data.get("photo_id")

    await callback.message.edit_text("⏳ Xabarnoma yuborilmoqda, iltimos kuting...")

    async with async_session() as session:
        if target_group == "completed":
            query = select(User.telegram_id).join(Profile, User.id == Profile.user_id).where(Profile.bio_summary.isnot(None))
        elif target_group == "incomplete":
            query = (
                select(User.telegram_id)
                .outerjoin(Profile, User.id == Profile.user_id)
                .where(or_(Profile.id.is_(None), Profile.bio_summary.is_(None)))
            )
        else:
            query = select(User.telegram_id)

        target_ids = (await session.execute(query)).scalars().all()

    success = 0
    failed = 0

    for tg_id in target_ids:
        try:
            if photo_id:
                await callback.bot.send_photo(chat_id=tg_id, photo=photo_id, caption=text, parse_mode=None)
            else:
                await callback.bot.send_message(chat_id=tg_id, text=text, parse_mode=None)
            success += 1
            await asyncio.sleep(0.04)
        except Exception:
            failed += 1

    await callback.message.answer(
        "🎉 **XABARNOMA TARQATILDI!**\n\n"
        f"• Muvaffaqiyatli yetkazildi: **{success} ta**\n"
        f"• Yetkazilmadi (bloklagan/o'chirilgan): **{failed} ta**",
        reply_markup=admin_main_keyboard(),
        parse_mode="Markdown"
    )


# ═══════════════════════════════════════════════════════════════════════
# EXCEL (.XLSX) EKSPORT
# ═══════════════════════════════════════════════════════════════════════

@router.message(Command("export"))
@router.callback_query(F.data == "admin_export")
async def admin_export_excel(event: Message | CallbackQuery):
    """Barcha a'zolar ma'lumotlarini to'liq Excel (.xlsx) faylida shakllantirib berish."""
    user_id = event.from_user.id
    if not await is_admin_user(user_id):
        return

    if isinstance(event, CallbackQuery):
        await event.answer("Excel shakllantirilmoqda...")
        reply_target = event.message
    else:
        reply_target = event

    async with async_session() as session:
        query = (
            select(User, Profile)
            .outerjoin(Profile, User.id == Profile.user_id)
            .order_by(User.id.asc())
        )
        rows = (await session.execute(query)).all()

    now_str = datetime.now().strftime("%Y_%m_%d_%H%M")
    fieldnames = [
        "ID",
        "Telegram ID",
        "Ism-Familiya",
        "Telegram Username",
        "Telefon raqam",
        "Anketa Holati",
        "Yoshi",
        "Kasbi / Lavozimi",
        "Kompaniya / Loyiha",
        "Faoliyat sohasi",
        "Ish tajribasi (yil)",
        "Eng katta yutug'i va natijasi",
        "Qidirayotgan sherigi",
        "Qaysi sohadan sherik izlamoqda",
        "Qidirilayotgan daraja",
        "O'zi nima taklif qila oladi",
        "Muhokama mavzulari",
        "Qisqa xulosa (Bio Summary)",
        "Ro'yxatdan o'tgan sana",
    ]

    export_rows = []
    for u, p in rows:
        created_at_str = u.created_at.strftime("%Y-%m-%d %H:%M") if hasattr(u, "created_at") and u.created_at else ""
        if p and p.bio_summary:
            status = "To'liq topshirilgan"
        elif u.phone:
            status = "Jarayonda (Telefon bergan)"
        else:
            status = "Faqat botga kirgan"

        export_rows.append([
            u.id,
            u.telegram_id,
            u.first_name,
            f"@{u.username}" if u.username else "",
            u.phone or "",
            status,
            p.age if p and p.age else "",
            p.role if p and p.role else "",
            p.company if p and p.company else "",
            p.industry if p and p.industry else "",
            p.experience_years if p and p.experience_years else "",
            p.achievements if p and hasattr(p, "achievements") and p.achievements else "",
            p.target_partner if p and p.target_partner else (p.current_goal if p else ""),
            p.target_industry if p and p.target_industry else "",
            p.target_seniority if p and hasattr(p, "target_seniority") and p.target_seniority else "",
            p.can_offer if p and p.can_offer else "",
            p.interests if p and p.interests else "",
            p.bio_summary if p and p.bio_summary else "",
            created_at_str,
        ])

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        has_openpyxl = True
    except ImportError:
        has_openpyxl = False

    if has_openpyxl:
        filename = f"teahouse_members_{now_str}.xlsx"
        filepath = os.path.join(os.getcwd(), filename)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Teahouse A'zolari"

        ws.append(fieldnames)

        # Style header
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        for col_idx in range(1, len(fieldnames) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border

        ws.row_dimensions[1].height = 28

        for row_data in export_rows:
            ws.append(row_data)
            row_idx = ws.max_row
            ws.row_dimensions[row_idx].height = 22
            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        wb.save(filepath)
    else:
        filename = f"teahouse_members_{now_str}.csv"
        filepath = os.path.join(os.getcwd(), filename)
        with open(filepath, mode="w", newline="", encoding="utf-8-sig") as f:
            f.write("sep=;\n")
            writer = csv.writer(f, delimiter=";")
            writer.writerow(fieldnames)
            for r in export_rows:
                writer.writerow(r)

    try:
        doc = FSInputFile(filepath, filename=filename)
        caption = (
            f"📊 **Teahouse a'zolari to'liq ro'yxati (Excel)**\n\n"
            f"👥 Jami a'zolar: {len(rows)} nafar\n"
            f"📅 Vaqt: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"
            "✅ Ushbu faylni to'g'ridan-to'g'ri Microsoft Excel yoki Google Sheets dasturlarida ochishingiz mumkin."
        )
        await reply_target.answer_document(document=doc, caption=caption, parse_mode="Markdown")
    finally:
        if os.path.exists(filepath):
            try:
                os.remove(filepath)
            except Exception:
                pass
